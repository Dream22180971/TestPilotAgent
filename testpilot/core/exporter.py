"""Export test outputs to files: Markdown, Excel, JSON."""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

STORE_TARGETS = ["test_strategy", "test_plan", "test_points", "test_cases", "test_script"]
TARGET_LABELS = {
    "analysis": "需求分析",
    "test_strategy": "测试策略",
    "test_plan": "测试计划",
    "test_points": "测试点",
    "test_cases": "测试用例",
    "test_script": "测试脚本",
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="101828", end_color="101828", fill_type="solid")


def _safe_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", name).strip() or "testpilot"


def export_to_files(
    project_name: str,
    outputs: dict[str, Any],
    out_dir: Path,
    formats: tuple[str, ...] = ("md", "xlsx"),
) -> list[Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = _safe_name(project_name)
    written: list[Path] = []

    if "md" in formats:
        path = out_dir / f"{base}.md"
        path.write_text(render_markdown(project_name, outputs), encoding="utf-8")
        written.append(path)

    if "xlsx" in formats:
        path = out_dir / f"{base}.xlsx"
        build_workbook(project_name, outputs).save(path)
        written.append(path)

    if "json" in formats:
        path = out_dir / f"{base}.json"
        path.write_text(json.dumps(outputs, ensure_ascii=False, indent=2), encoding="utf-8")
        written.append(path)

    return written


def export_outputs(db, project_id: str, export_format: str, targets: list[str]) -> dict:
    """Legacy db-backed export used by the FastAPI layer."""
    from testpilot.core.db_store import DatabaseStore

    store = DatabaseStore()
    outputs: dict[str, Any] = {}
    for t in targets:
        if t in STORE_TARGETS:
            out = store.get_output(db, project_id, t)
            if out:
                outputs[t] = out["content"]

    if export_format == "json":
        return {"project_id": project_id, "format": "json", "status": "ready", "data": outputs}
    if export_format in ("markdown", "md"):
        return {
            "project_id": project_id,
            "format": "markdown",
            "status": "ready",
            "data": render_markdown(project_id, outputs),
        }
    if export_format in ("xlsx", "excel"):
        buf = io.BytesIO()
        build_workbook(project_id, outputs).save(buf)
        buf.seek(0)
        return {
            "project_id": project_id,
            "format": "xlsx",
            "status": "ready",
            "data": buf.getvalue().hex(),
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }
    return {"project_id": project_id, "format": export_format, "targets": targets, "status": "unsupported_format"}


def render_markdown(project_name: str, outputs: dict[str, Any]) -> str:
    lines = [f"# TestPilot 测试资产 · {project_name}", ""]
    for target, label in TARGET_LABELS.items():
        content = outputs.get(target)
        if content is None:
            continue
        lines.append(f"## {label}")
        lines.append("")
        if target == "test_cases":
            lines.extend(_cases_markdown(content))
        else:
            lines.append(_dict_to_md(content))
        lines.append("")
    return "\n".join(lines)


def _cases_markdown(content: dict) -> list[str]:
    lines = []
    for case in content.get("cases", []):
        prio = {"高": "P0", "中": "P1", "低": "P2"}.get(case.get("priority", "中"), case.get("priority", ""))
        lines.append(f"### {case.get('id', '')} `[{prio}]` {case.get('title', '')}")
        lines.append(f"- 模块：{case.get('module', '')} ｜ 场景：{case.get('scenario', '')}")
        pre = case.get("preconditions", [])
        if pre:
            lines.append(f"- 前置条件：{'; '.join(str(p) for p in pre)}")
        for i, step in enumerate(case.get("steps", []), start=1):
            lines.append(f"{i}. {step}")
        for exp in case.get("expected_results", []):
            lines.append(f"- 预期：{exp}")
        lines.append("")
    return lines


def build_workbook(project_name: str, outputs: dict[str, Any]) -> Workbook:
    wb = Workbook()
    for target, content in outputs.items():
        if target not in STORE_TARGETS or content is None:
            continue
        label = TARGET_LABELS.get(target, target)
        ws = wb.create_sheet(title=label[:31])
        if target == "test_points":
            rows = content.get("rows", [])
            if rows:
                _write_rows(ws, list(rows[0].keys()), rows)
        elif target == "test_cases":
            cases = content.get("cases", [])
            if cases:
                headers = ["id", "module", "scenario", "title", "preconditions", "steps", "expected_results", "priority"]
                _write_rows(ws, headers, cases)
        else:
            _write_dict_sheet(ws, content)
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    return wb


def _write_rows(ws, headers: list[str], rows: list[dict]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            val = row.get(header, "")
            if isinstance(val, list):
                val = "; ".join(str(v) for v in val)
            ws.cell(row=row_idx, column=col_idx, value=val)
    ws.freeze_panes = "A2"


def _write_dict_sheet(ws, content: dict) -> None:
    items = list(content.items())
    if not items:
        return
    for col_idx, (key, val) in enumerate(items, start=1):
        cell = ws.cell(row=1, column=col_idx, value=key)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        if isinstance(val, list):
            val = "\n".join(str(v) for v in val)
        elif isinstance(val, dict):
            val = json.dumps(val, ensure_ascii=False)
        ws.cell(row=2, column=col_idx, value=str(val))


def _dict_to_md(data: Any, indent: int = 0) -> str:
    prefix = "  " * indent
    if isinstance(data, dict):
        lines = []
        for key, val in data.items():
            if isinstance(val, (dict, list)):
                lines.append(f"{prefix}- **{key}**:")
                lines.append(_dict_to_md(val, indent + 1))
            else:
                lines.append(f"{prefix}- **{key}**: {val}")
        return "\n".join(lines)
    elif isinstance(data, list):
        lines = []
        for item in data:
            if isinstance(item, (dict, list)):
                lines.append(_dict_to_md(item, indent))
            else:
                lines.append(f"{prefix}- {item}")
        return "\n".join(lines)
    return f"{prefix}{data}"
